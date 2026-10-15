"""演示新版：数字页 / 横条对比、不做章节分隔页、风格参数（默认瑞士网格 + 一句话描述）。"""

import os
import tempfile
import unittest
from unittest import mock

from apps.notes2insight import deck, deck_theme, llm


SOURCES = [{"idx": 1, "title": "t1", "date": "2026-01-01", "folder": "f", "rel": "a.md"},
           {"idx": 2, "title": "t2", "date": "", "folder": "f", "rel": "b.md"}]


def _deck(**kw):
    d = {"title": "报告", "subtitle": "副标题", "meta": {}, "sources": SOURCES, "slides": [
        {"kind": "points", "section": "开篇", "title": "一", "lead": "", "note": "n",
         "bullets": [{"text": "要点", "cites": [1]}]},
        {"kind": "stat", "section": "数字", "title": "二", "lead": "", "note": "",
         "items": [{"value": "5GW", "label": "规模", "cites": [1]}, {"value": "3-4×", "label": "增速", "cites": []}]},
        {"kind": "bars", "section": "价格", "title": "三", "lead": "", "note": "", "unit": "$M/MW",
         "items": [{"label": "现在", "value": 12.5, "display": "$10-15M", "cites": [2]},
                   {"label": "目标", "value": 100, "display": "$100M", "cites": []}]},
        {"kind": "timeline", "section": "情景", "title": "四", "lead": "", "note": "",
         "events": [{"date": "2026", "text": "x", "cites": []}, {"date": "2027", "text": "y", "cites": [1]}]},
    ]}
    d.update(kw)
    return d


class NormalizeV2Tests(unittest.TestCase):
    def test_数字页和横条对比_数字不全的降级或丢掉_章节分隔页不要(self):
        slides = deck.normalize({"slides": [
            {"kind": "section", "title": "第 1 章"},
            {"kind": "stat", "title": "数字", "items": [{"value": "5GW", "label": "规模", "cites": [1, 9]},
                                                        {"value": "", "label": "空的"}]},
            {"kind": "stat", "title": "全空", "items": [{"label": "没数字"}]},
            {"kind": "bars", "title": "对比", "unit": "$M",
             "items": [{"label": "A", "value": "12.5", "display": "$12.5M"}, {"label": "B", "value": 100}]},
            {"kind": "bars", "title": "凑不成", "items": [{"label": "A", "value": "很多", "display": "很多"},
                                                         {"label": "B", "value": 3, "display": "3 家"}]},
        ]}, SOURCES)
        self.assertEqual([(s["kind"], s["title"]) for s in slides],
                         [("stat", "数字"), ("bars", "对比"), ("points", "凑不成")])
        self.assertEqual(slides[0]["items"], [{"value": "5GW", "label": "规模", "cites": [1]}])
        self.assertEqual([i["value"] for i in slides[1]["items"]], [12.5, 100.0])
        self.assertEqual(slides[1]["items"][1]["display"], "100")
        self.assertEqual(slides[2]["bullets"], [{"text": "B：3 家", "cites": []}])


class ThemeTests(unittest.TestCase):
    def test_默认是瑞士网格_不合规的值退回默认_对比度不够整组颜色退回(self):
        t = deck_theme.clean_theme({"accent": "#d9480f", "font": "comic", "radius": 99, "title_weight": 650,
                                    "evil": "x"})
        self.assertEqual((t["accent"], t["font"], t["radius"], t["title_weight"]), ("#d9480f", "sans", 0, 600))
        self.assertNotIn("evil", t)
        dark = deck_theme.clean_theme({"bg": "#111111", "ink": "#f5f5f5", "dim": "#a0a0a0", "accent": "#ff8a3d"})
        self.assertEqual(dark["bg"], "#111111")
        unreadable = deck_theme.clean_theme({"bg": "#ffffff", "ink": "#eeeeee"})
        self.assertEqual(unreadable["ink"], deck_theme.DEFAULT_THEME["ink"])

    def test_一句话描述转成风格参数_空描述不调模型(self):
        with mock.patch.object(llm, "complete") as m:
            self.assertEqual(deck_theme.theme_from_description("", backend="api")["name"], "瑞士网格")
            m.assert_not_called()
        raw = '好的：{"name": "深色", "bg": "#121212", "ink": "#f0f0f0", "dim": "#a8a8a8", "accent": "#ff8a3d", "display_font": "serif"}'
        with mock.patch.object(llm, "complete", return_value=raw):
            t = deck_theme.theme_from_description("深色背景、橙色点缀、标题衬线", backend="api")
        self.assertEqual((t["bg"], t["accent"], t["display_font"], t["description"]),
                         ("#121212", "#ff8a3d", "serif", "深色背景、橙色点缀、标题衬线"))


class RenderV2Tests(unittest.TestCase):
    def test_html带风格变量_风格存进去读得回来_老演示的分隔页和配图去掉(self):
        theme = deck_theme.clean_theme({"accent": "#d9480f"})
        html = deck.render_html(_deck(theme=theme))
        self.assertIn("--accent:#d9480f", html)
        self.assertNotIn("__THEME_VARS__", html)
        with tempfile.TemporaryDirectory() as d:
            p = os.path.join(d, "r.deck.html")
            with open(p, "w", encoding="utf-8") as f:
                f.write(html)
            back = deck.load_deck_from_html(p)
            self.assertEqual(back["theme"]["accent"], "#d9480f")
            self.assertEqual([s["kind"] for s in back["slides"]], ["points", "stat", "bars", "timeline"])
            old = html.replace('"kind": "points"', '"kind": "section"', 1).replace('"theme":', '"x_theme":')
            with open(p, "w", encoding="utf-8") as f:
                f.write(old)
            back = deck.load_deck_from_html(p)
        self.assertEqual(back["theme"]["name"], "瑞士网格")        # 老演示没存风格 → 默认
        self.assertNotIn("section", [s["kind"] for s in back["slides"]])

    def test_pptx_新页面类型都能导出(self):
        with tempfile.TemporaryDirectory() as d:
            out = os.path.join(d, "r.pptx")
            deck.render_pptx(_deck(theme={"accent": "#d9480f"}), out)
            from pptx import Presentation
            prs = Presentation(out)
        self.assertEqual(len(prs.slides), 6)    # 封面 + 4 页 + 来源
        texts = " ".join(sh.text_frame.text for sh in prs.slides[2].shapes if sh.has_text_frame)
        self.assertIn("5GW", texts)


if __name__ == "__main__":
    unittest.main()
